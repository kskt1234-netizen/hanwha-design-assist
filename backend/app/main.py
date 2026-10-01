from contextlib import asynccontextmanager
from io import BytesIO
from typing import Literal

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import PlainTextResponse
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import repository
from app.db import get_session, init_db
from app.generators.ddl import generate_ddl
from app.generators.mermaid import generate_mermaid
from app.generators.trace import build_trace
from app.llm import LLMNotConfigured
from app.models import STAGES, ArtifactVersion, GenerationRun, Project, ReviewIssueRow, SourceDocument
from app.pipeline.stages import GENERATORS, StageError, run_review
from app.rag.retriever import available_domains
from app.schemas import DataModel, DesignState, FeatureSpec, RequirementAnalysis, TraceRow

Stage = Literal["requirements", "features", "data_model"]


@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()
    yield


app = FastAPI(title="AI 설계 지원 서비스", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


# ---------- 응답 모델 ----------

class ProjectIn(BaseModel):
    name: str
    domain: str


class ProjectOut(BaseModel):
    id: int
    name: str
    domain: str
    stage_status: dict[str, str]
    document_count: int


class DocumentOut(BaseModel):
    id: int
    filename: str
    length: int


class TextDocumentIn(BaseModel):
    filename: str = "requirements.txt"
    text: str


class IssueOut(BaseModel):
    id: int
    rule_code: str
    severity: str
    target_type: str
    target_code: str
    message: str
    status: str


class IssueStatusIn(BaseModel):
    status: Literal["OPEN", "RESOLVED", "IGNORED"]


class VersionOut(BaseModel):
    id: int
    stage: str
    version_no: int
    created_at: str


def _project(session: Session, project_id: int) -> Project:
    project = session.get(Project, project_id)
    if not project:
        raise HTTPException(404, "프로젝트가 없습니다.")
    return project


def _project_out(p: Project) -> ProjectOut:
    return ProjectOut(id=p.id, name=p.name, domain=p.domain, stage_status=p.stage_status, document_count=len(p.documents))


# ---------- 도메인 / 프로젝트 ----------

@app.get("/api/domains")
def list_domains() -> list[dict]:
    return available_domains()


@app.post("/api/projects", status_code=201)
def create_project(body: ProjectIn, session: Session = Depends(get_session)) -> ProjectOut:
    if body.domain not in {d["code"] for d in available_domains()}:
        raise HTTPException(400, f"알 수 없는 도메인: {body.domain}")
    p = Project(name=body.name, domain=body.domain, stage_status={s: "EMPTY" for s in STAGES})
    session.add(p)
    session.commit()
    return _project_out(p)


@app.get("/api/projects")
def list_projects(session: Session = Depends(get_session)) -> list[ProjectOut]:
    return [_project_out(p) for p in session.scalars(select(Project).order_by(Project.id.desc()))]


@app.get("/api/projects/{project_id}")
def get_project(project_id: int, session: Session = Depends(get_session)) -> ProjectOut:
    return _project_out(_project(session, project_id))


# ---------- 요구사항 문서 ----------

def _extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    if name.endswith(".docx"):
        from docx import Document

        return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
    return data.decode("utf-8", errors="replace")


@app.post("/api/projects/{project_id}/documents", status_code=201)
async def upload_document(
    project_id: int, file: UploadFile = File(...), session: Session = Depends(get_session)
) -> DocumentOut:
    project = _project(session, project_id)
    raw = _extract_text(file.filename or "upload.txt", await file.read())
    doc = SourceDocument(project_id=project.id, filename=file.filename or "upload.txt", raw_text=raw)
    session.add(doc)
    session.commit()
    return DocumentOut(id=doc.id, filename=doc.filename, length=len(raw))


@app.post("/api/projects/{project_id}/documents/text", status_code=201)
def add_text_document(project_id: int, body: TextDocumentIn, session: Session = Depends(get_session)) -> DocumentOut:
    project = _project(session, project_id)
    doc = SourceDocument(project_id=project.id, filename=body.filename, raw_text=body.text)
    session.add(doc)
    session.commit()
    return DocumentOut(id=doc.id, filename=doc.filename, length=len(body.text))


@app.get("/api/projects/{project_id}/documents")
def list_documents(project_id: int, session: Session = Depends(get_session)) -> list[DocumentOut]:
    return [DocumentOut(id=d.id, filename=d.filename, length=len(d.raw_text)) for d in _project(session, project_id).documents]


@app.delete("/api/projects/{project_id}/documents/{document_id}", status_code=204)
def delete_document(project_id: int, document_id: int, session: Session = Depends(get_session)) -> None:
    doc = session.get(SourceDocument, document_id)
    if not doc or doc.project_id != project_id:
        raise HTTPException(404, "문서가 없습니다.")
    session.delete(doc)
    session.commit()


# ---------- 설계 단계: 생성 / 편집 / 승인 ----------

@app.post("/api/projects/{project_id}/stages/{stage}/generate")
def generate_stage(project_id: int, stage: Stage, session: Session = Depends(get_session)) -> DesignState:
    project = _project(session, project_id)
    try:
        GENERATORS[stage](session, project)
    except (StageError, LLMNotConfigured) as exc:
        raise HTTPException(400, str(exc)) from exc
    return repository.load_state(session, project.id)


@app.get("/api/projects/{project_id}/design")
def get_design(project_id: int, session: Session = Depends(get_session)) -> DesignState:
    return repository.load_state(session, _project(session, project_id).id)


def _mark_edited(project: Project, stage: str) -> None:
    project.stage_status = {**project.stage_status, stage: "DRAFT"}


@app.put("/api/projects/{project_id}/design/requirements")
def save_requirements(project_id: int, body: RequirementAnalysis, session: Session = Depends(get_session)) -> DesignState:
    project = _project(session, project_id)
    repository.save_requirements(session, project.id, body.requirements)
    _mark_edited(project, "requirements")
    session.commit()
    return repository.load_state(session, project.id)


@app.put("/api/projects/{project_id}/design/features")
def save_features(project_id: int, body: FeatureSpec, session: Session = Depends(get_session)) -> DesignState:
    project = _project(session, project_id)
    repository.save_features(session, project.id, body.features)
    _mark_edited(project, "features")
    session.commit()
    return repository.load_state(session, project.id)


@app.put("/api/projects/{project_id}/design/data_model")
def save_data_model(project_id: int, body: DataModel, session: Session = Depends(get_session)) -> DesignState:
    project = _project(session, project_id)
    repository.save_data_model(session, project.id, body.entities, body.relationships)
    _mark_edited(project, "data_model")
    session.commit()
    return repository.load_state(session, project.id)


@app.post("/api/projects/{project_id}/stages/{stage}/approve")
def approve_stage(project_id: int, stage: Stage, session: Session = Depends(get_session)) -> VersionOut:
    project = _project(session, project_id)
    if project.stage_status.get(stage) == "EMPTY":
        raise HTTPException(400, "승인할 초안이 없습니다.")
    state = repository.load_state(session, project.id)
    last = session.scalar(
        select(func.max(ArtifactVersion.version_no)).where(
            ArtifactVersion.project_id == project.id, ArtifactVersion.stage == stage
        )
    )
    version = ArtifactVersion(
        project_id=project.id, stage=stage, version_no=(last or 0) + 1, snapshot=state.model_dump(mode="json")
    )
    session.add(version)
    project.stage_status = {**project.stage_status, stage: "APPROVED"}
    session.commit()
    return VersionOut(id=version.id, stage=stage, version_no=version.version_no, created_at=version.created_at.isoformat())


@app.get("/api/projects/{project_id}/versions")
def list_versions(project_id: int, session: Session = Depends(get_session)) -> list[VersionOut]:
    rows = session.scalars(
        select(ArtifactVersion)
        .where(ArtifactVersion.project_id == _project(session, project_id).id)
        .order_by(ArtifactVersion.id.desc())
    )
    return [VersionOut(id=v.id, stage=v.stage, version_no=v.version_no, created_at=v.created_at.isoformat()) for v in rows]


@app.get("/api/projects/{project_id}/runs")
def list_runs(project_id: int, session: Session = Depends(get_session)) -> list[dict]:
    rows = session.scalars(
        select(GenerationRun).where(GenerationRun.project_id == project_id).order_by(GenerationRun.id.desc()).limit(50)
    )
    return [
        {"stage": r.stage, "provider": r.provider, "model": r.model, "status": r.status,
         "retrieved_chunk_ids": r.retrieved_chunk_ids, "started_at": r.started_at.isoformat()}
        for r in rows
    ]


# ---------- 추적표 / 검토 / 내보내기 ----------

@app.get("/api/projects/{project_id}/trace")
def get_trace(project_id: int, session: Session = Depends(get_session)) -> list[TraceRow]:
    return build_trace(repository.load_state(session, _project(session, project_id).id))


@app.post("/api/projects/{project_id}/review")
def review(project_id: int, session: Session = Depends(get_session)) -> list[IssueOut]:
    project = _project(session, project_id)
    try:
        run_review(session, project)
    except LLMNotConfigured as exc:
        raise HTTPException(400, str(exc)) from exc
    return list_issues(project_id, session)


@app.get("/api/projects/{project_id}/issues")
def list_issues(project_id: int, session: Session = Depends(get_session)) -> list[IssueOut]:
    rows = session.scalars(
        select(ReviewIssueRow).where(ReviewIssueRow.project_id == project_id).order_by(ReviewIssueRow.id)
    )
    return [
        IssueOut(id=r.id, rule_code=r.rule_code, severity=r.severity, target_type=r.target_type,
                 target_code=r.target_code, message=r.message, status=r.status)
        for r in rows
    ]


@app.patch("/api/projects/{project_id}/issues/{issue_id}")
def update_issue(project_id: int, issue_id: int, body: IssueStatusIn, session: Session = Depends(get_session)) -> IssueOut:
    row = session.get(ReviewIssueRow, issue_id)
    if not row or row.project_id != project_id:
        raise HTTPException(404, "검토 항목이 없습니다.")
    row.status = body.status
    session.commit()
    return IssueOut(id=row.id, rule_code=row.rule_code, severity=row.severity, target_type=row.target_type,
                    target_code=row.target_code, message=row.message, status=row.status)


@app.get("/api/projects/{project_id}/export/ddl", response_class=PlainTextResponse)
def export_ddl(project_id: int, session: Session = Depends(get_session)) -> str:
    return generate_ddl(repository.load_state(session, _project(session, project_id).id))


@app.get("/api/projects/{project_id}/export/mermaid", response_class=PlainTextResponse)
def export_mermaid(project_id: int, session: Session = Depends(get_session)) -> str:
    return generate_mermaid(repository.load_state(session, _project(session, project_id).id))


@app.get("/api/health")
def health() -> dict:
    return {"status": "ok"}
